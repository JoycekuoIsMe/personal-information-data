from openpyxl import *
from openpyxl.drawing.image import Image
import docx
from openpyxl import Workbook
import docx

class PersonInfo:
    def __init__(self, fullName, id_number):
        self.fullName = fullName
        self.id_number = id_number

    def get_name(self):
        return self.fullName.split(" ")

    def get_first_name(self):
        nameParts = self.get_name()
        if len(nameParts) >= 1:
            return nameParts[0]
        return ""

    def get_last_name(self):
        nameParts = self.get_name()
        if len(nameParts) >= 2:
            return nameParts[1]
        return ""
    
    def get_id_number(self):
        return self.id_number.replace(',', '')

    def get_gender(self):
        id_code = self.id_number[1]
        if id_code == '1':
            return "Male"
        elif id_code == '2':
            return "Female"
        return "Unknown"

    def get_hometown(self):
        hometown_map = {
            'A': '台北市',
            'B': '台中市',
            'C': '基隆市',
            'D': '台南市',
            'E': '高雄市',
            'F': '新北市',
            'G': '宜蘭縣',
            'H': '桃園市',
            'J': '新竹縣',
            'K': '苗栗縣',
            'L': '台中縣',
            'M': '南投縣',
            'N': '彰化縣',
            'P': '雲林縣',
            'Q': '嘉義縣',
            'R': '台南縣',
            'S': '高雄縣',
            'T': '屏東縣',
            'U': '花蓮縣',
            'V': '台東縣',
            'W': '金門縣',
            'X': '澎湖縣',
            'Y': '陽明山',
            'Z': '連江縣',
            'I': '嘉義市',
            'O': '新竹市'
        }
        hometownCode = self.id_number[0]
        return hometown_map.get(hometownCode, '')


# 開啟word文檔
doc = docx.Document("./身分資料文件.docx")

# 新增Excel工作簿和工作表
workbook = Workbook()
sheet = workbook.active
sheet.title = "個人資訊"

# 添加表頭
sheet.append(["姓", "名", "性別","身份證", "戶籍地" ])

# 讀取word文檔段落
for paragraph in doc.paragraphs:
    # 將段落依','切割
    data = paragraph.text.split(", ")

    # 跳過不符合預期的數據格式 " "
    if len(data) != 2:
        continue

    # 創建PersonInfo object
    personInfo = PersonInfo(data[0], data[1])

    # 獲取個人訊息
    lastName = personInfo.get_last_name()
    firstName = personInfo.get_first_name()
    gender = personInfo.get_gender()
    hometown = personInfo.get_hometown()
    idNum = personInfo.get_id_number()

    # 寫入Excel文件
    sheet.append([lastName,firstName, gender,idNum , hometown])

# 調整列寬
for col in sheet.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        if len(str(cell.value)) > max_length:
                max_length = len(cell.value)
    adjusted_width = (max_length + 2)
    sheet.column_dimensions[column].width = adjusted_width

# 保存Excel文件
workbook.save("./Personal Information.xlsx")
